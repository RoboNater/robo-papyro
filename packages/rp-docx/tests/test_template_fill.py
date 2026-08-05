"""``{{ placeholder }}`` substitution without docxtpl (spec section 8)."""

from __future__ import annotations

import pytest

from rp_docx import ooxml
from rp_docx.docx import read, template
from rp_docx.errors import PlaceholderError, TemplateError


class TestSyntax:
    def test_a_plain_key(self):
        assert template.PLACEHOLDER.findall("{{ name }}") == ["name"]

    def test_a_dotted_key(self):
        assert template.PLACEHOLDER.findall("{{ client.name }}") == ["client.name"]

    def test_whitespace_is_optional(self):
        assert template.PLACEHOLDER.findall("{{name}} {{  name  }}") == ["name", "name"]

    def test_expressions_are_not_syntax(self):
        """No expression evaluation and no Jinja: a template is data, and
        rendering one must not be able to run anything."""
        for hostile in (
            "{{ 1 + 1 }}",
            "{{ os.system('rm -rf /') }}",
            "{{ obj['key'] }}",
            "{% for x in y %}",
            "{{ a|filter }}",
        ):
            assert template.PLACEHOLDER.findall(hostile) == []

    def test_a_call_is_not_a_key(self):
        assert template.PLACEHOLDER.findall("{{ name() }}") == []


class TestResolution:
    def test_a_flat_key(self):
        assert template.resolve({"name": "Ada"}, "name") == "Ada"

    def test_a_nested_key(self):
        assert template.resolve({"client": {"name": "Ada"}}, "client.name") == "Ada"

    def test_a_missing_key_is_none_not_the_string_none(self):
        """So an unresolved placeholder is reported as unresolved rather than
        rendered as "None"."""
        assert template.resolve({}, "name") is None
        assert template.resolve({"client": {}}, "client.name") is None

    def test_descending_into_a_non_mapping_is_a_miss(self):
        assert template.resolve({"client": "Ada"}, "client.name") is None

    def test_non_string_values_are_stringified(self):
        assert template.resolve({"n": 42}, "n") == "42"

    def test_an_explicit_null_becomes_an_empty_string(self):
        """A key that is present with a null value was supplied — it just has no
        text — which is different from one that is missing."""
        assert template.resolve({"n": None}, "n") == ""


class TestDiscovery:
    def test_placeholders_are_found_across_runs(self, split_runs_docx):
        assert set(template.find_placeholders(split_runs_docx)) == {
            "client.name",
            "amount",
            "city",
        }

    def test_placeholders_in_headers_are_found(self, split_runs_docx):
        """Found before someone discovers them in a printed document."""
        assert "city" in template.find_placeholders(split_runs_docx)

    def test_a_document_with_no_placeholders(self, simple_docx):
        assert template.find_placeholders(simple_docx) == []


class TestFill:
    CONTEXT = {"client": {"name": "Ada"}, "amount": "£40", "city": "Bath"}

    def test_filling_replaces_everywhere(self, split_runs_docx, tmp_path):
        result = template.fill_template(split_runs_docx, self.CONTEXT, tmp_path / "filled.docx")
        assert result.unresolved == []
        text = " ".join(p.text for p in read.get_text(result.output))
        assert "Dear Ada, welcome." in text
        assert "Total: £40 due" in text

    def test_filling_reports_what_it_filled(self, split_runs_docx, tmp_path):
        result = template.fill_template(split_runs_docx, self.CONTEXT, tmp_path / "filled.docx")
        assert result.filled == {"client.name": "Ada", "amount": "£40", "city": "Bath"}

    def test_strict_refuses_a_half_filled_document(self, split_runs_docx, tmp_path):
        """A contract with {{ client.name }} still in it is worse than no
        document at all."""
        with pytest.raises(PlaceholderError) as exc:
            template.fill_template(split_runs_docx, {"amount": "£40"}, tmp_path / "f.docx")
        assert "client.name" in str(exc.value)
        assert "city" in str(exc.value)

    def test_strict_failure_is_an_input_error(self, split_runs_docx, tmp_path):
        with pytest.raises(PlaceholderError) as exc:
            template.fill_template(split_runs_docx, {}, tmp_path / "f.docx")
        assert exc.value.exit_code == 1

    def test_non_strict_leaves_them_and_reports_them(self, split_runs_docx, tmp_path):
        result = template.fill_template(
            split_runs_docx, {"amount": "£40"}, tmp_path / "f.docx", strict=False
        )
        assert sorted(result.unresolved) == ["city", "client.name"]
        assert "{{ client.name }}" in " ".join(p.text for p in read.get_text(result.output))

    def test_the_template_is_never_modified(self, split_runs_docx, tmp_path):
        before = split_runs_docx.read_bytes()
        template.fill_template(split_runs_docx, self.CONTEXT, tmp_path / "f.docx")
        assert split_runs_docx.read_bytes() == before

    def test_a_bare_template_name_resolves(self, templates_env, minimal_template, tmp_path):
        result = template.fill_template("minimal", {}, tmp_path / "f.docx")
        assert result.output.is_file()

    def test_an_unknown_template_name_is_reported(self, templates_env, tmp_path):
        with pytest.raises(TemplateError, match="No template called"):
            template.fill_template("nonexistent", {}, tmp_path / "f.docx")

    def test_filling_a_dotx_writes_a_document(self, minimal_template, tmp_path):
        """Copying the package would carry the template content type across, and
        Word opens that as a template — silently creating an untitled copy
        instead of the file the user asked for."""
        result = template.fill_template(minimal_template, {}, tmp_path / "filled.docx")
        assert not ooxml.is_template(result.output)

    def test_filling_into_a_dotx_keeps_it_a_template(self, simple_docx, tmp_path):
        result = template.fill_template(simple_docx, {}, tmp_path / "filled.dotx")
        assert ooxml.is_template(result.output)

    def test_every_spelling_of_a_key_is_filled(self, tmp_path):
        """ "{{name}}" and "{{ name }}" are the same placeholder to everyone
        except a literal string match."""
        import docx

        document = docx.Document()
        document.add_paragraph("{{name}} and {{ name }} and {{  name  }}")
        source = tmp_path / "spellings.docx"
        document.save(str(source))
        result = template.fill_template(source, {"name": "Ada"}, tmp_path / "f.docx")
        assert read.get_text(result.output)[0].text == "Ada and Ada and Ada"
